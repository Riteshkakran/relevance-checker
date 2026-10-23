"""Turns raw bytes of ANY supported format into (kind, plain text, outgoing links).

HOW TO ADD A NEW FORMAT (this is the scalability point):
  1. Add its extension to SUPPORTED_EXTS.
  2. Teach detect_kind() to recognise it.
  3. Write a function decorated with @extractor("<kind>") that returns (text, links).
Nothing else in the project needs to change.
"""
import io
import os
import re
import zipfile
from urllib.parse import urljoin, urlparse

URL_RE = re.compile(r"https?://[^\s<>\"'\)\]]+")

SUPPORTED_EXTS = {
    ".pdf", ".docx", ".pptx", ".xlsx",
    ".html", ".htm",
    ".txt", ".md", ".csv", ".tsv", ".json", ".log",
}
TEXT_EXTS = {".txt", ".md", ".csv", ".tsv", ".json", ".log"}
OLE_MAGIC = bytes.fromhex("D0CF11E0A1B11AE1")  # old .doc / .xls / .ppt

EXTRACTORS = {}


def extractor(kind):
    def wrap(fn):
        EXTRACTORS[kind] = fn
        return fn
    return wrap


def _ext(name: str) -> str:
    path = urlparse(name).path if "://" in name else name
    return os.path.splitext(path)[1].lower()


def _merge_links(text: str, links):
    found = list(links) + [u.rstrip(".,;:") for u in URL_RE.findall(text)]
    return list(dict.fromkeys(found))


def detect_kind(content: bytes, ctype: str, name: str) -> str:
    ext = _ext(name)
    ctype = (ctype or "").lower()
    if content[:5] == b"%PDF-" or "pdf" in ctype:
        return "pdf"
    if content[:4] == b"PK\x03\x04":
        try:
            names = set(zipfile.ZipFile(io.BytesIO(content)).namelist())
        except zipfile.BadZipFile:
            raise ValueError("This file looks damaged.")
        if "word/document.xml" in names:
            return "docx"
        if "ppt/presentation.xml" in names:
            return "pptx"
        if "xl/workbook.xml" in names:
            return "xlsx"
    if content[:8] == OLE_MAGIC:
        raise ValueError("Old Office files (.doc, .xls, .ppt) aren't supported. Save as .docx, .xlsx or .pptx and try again.")
    head = content[:300].lower().lstrip()
    if "html" in ctype or ext in (".html", ".htm") or head.startswith((b"<!doctype html", b"<html")):
        return "html"
    if ext in TEXT_EXTS or ctype.startswith("text/") or "json" in ctype:
        return "text"
    raise ValueError("Unsupported file type. Supported: PDF, Word, PowerPoint, Excel, HTML, TXT, MD, CSV, JSON.")


def extract(content: bytes, ctype: str, name: str):
    kind = detect_kind(content, ctype, name)
    text, links = EXTRACTORS[kind](content, name)
    return kind, text, links


# ---------------------------------------------------------------- extractors
@extractor("pdf")
def _pdf(content, base):
    import fitz  # PyMuPDF
    doc = fitz.open(stream=content, filetype="pdf")
    parts, links = [], []
    for page in doc:
        parts.append(page.get_text())
        links += [l["uri"] for l in page.get_links() if l.get("uri")]
    text = "\n".join(parts)
    return text, _merge_links(text, links)


@extractor("html")
def _html(content, base):
    from bs4 import BeautifulSoup
    soup = BeautifulSoup(content.decode("utf-8", errors="ignore"), "html.parser")
    links = [urljoin(base, a["href"]) for a in soup.find_all("a", href=True)]
    for tag in soup(["script", "style", "noscript"]):
        tag.decompose()
    return soup.get_text(" ", strip=True), list(dict.fromkeys(links))


@extractor("docx")
def _docx(content, base):
    import docx
    d = docx.Document(io.BytesIO(content))
    parts = [p.text for p in d.paragraphs]
    for table in d.tables:
        for row in table.rows:
            parts += [cell.text for cell in row.cells]
    links = [r.target_ref for r in d.part.rels.values()
             if r.reltype.endswith("/hyperlink") and r.is_external]
    text = "\n".join(parts)
    return text, _merge_links(text, links)


@extractor("pptx")
def _pptx(content, base):
    from pptx import Presentation
    prs = Presentation(io.BytesIO(content))
    parts, links = [], []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    parts.append("".join(r.text for r in para.runs))
                    links += [r.hyperlink.address for r in para.runs if r.hyperlink.address]
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    parts += [c.text for c in row.cells]
        if slide.has_notes_slide:
            parts.append(slide.notes_slide.notes_text_frame.text)
    text = "\n".join(parts)
    return text, _merge_links(text, links)


@extractor("xlsx")
def _xlsx(content, base):
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    parts, cells = [], 0
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            vals = [str(v) for v in row if v is not None]
            if vals:
                parts.append(" ".join(vals))
                cells += len(vals)
            if cells > 200_000:          # safety cap for huge sheets
                break
    text = "\n".join(parts)
    return text, _merge_links(text, [])


@extractor("text")
def _text(content, base):
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = content.decode("latin-1")
    return text, _merge_links(text, [])
